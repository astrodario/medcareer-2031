import openpyxl
import pandas as pd
import re

wb = openpyxl.load_workbook("data/raw/mur_ssm/riparto_SSM_22_23.xlsx", data_only=True)

uni_to_region = {
    'BARI': 'Puglia',
    'BOLOGNA': 'Emilia-Romagna',
    'BRESCIA': 'Lombardia',
    'CAGLIARI': 'Sardegna',
    'CAMPANIA - "L. VANVITELLI"': 'Campania',
    'CATANIA': 'Sicilia',
    'CATANZARO': 'Calabria',
    'CHIETI-PESCARA': 'Abruzzo',
    'Cattolica del Sacro Cuore': 'Lazio',
    'FERRARA': 'Emilia-Romagna',
    'FIRENZE': 'Toscana',
    'FOGGIA': 'Puglia',
    'GENOVA': 'Liguria',
    'HUMANITAS University': 'Lombardia',
    'INSUBRIA': 'Lombardia',
    "L'AQUILA": 'Abruzzo',
    'MESSINA': 'Sicilia',
    'MILANO': 'Lombardia',
    'MILANO-BICOCCA': 'Lombardia',
    'MODENA e REGGIO EMILIA': 'Emilia-Romagna',
    'MOLISE': 'Molise',
    'Napoli Federico II': 'Campania',
    'PADOVA': 'Veneto',
    'PALERMO': 'Sicilia',
    'PARMA': 'Emilia-Romagna',
    'PAVIA': 'Lombardia',
    'PERUGIA': 'Umbria',
    'PIEMONTE ORIENTALE': 'Piemonte',
    'PISA': 'Toscana',
    'Politecnica delle MARCHE': 'Marche',
    'ROMA "La Sapienza" Fac. M-O/F-M': 'Lazio',
    'ROMA "La Sapienza" Fac. M-P': 'Lazio',
    'ROMA "Tor Vergata"': 'Lazio',
    'S. Raffaele MILANO': 'Lombardia',
    'SALERNO': 'Campania',
    'SASSARI': 'Sardegna',
    'SIENA': 'Toscana',
    'TORINO': 'Piemonte',
    'TRIESTE': 'Friuli-Venezia Giulia',
    'UDINE': 'Friuli-Venezia Giulia',
    'UniCamillus - Saint Camillus International U': 'Lazio',
    'Univ. "Campus Bio-Medico" di ROMA': 'Lazio',
    'VERONA': 'Veneto',
    'della CALABRIA': 'Calabria'
}

all_regional_records = []
summary_records = []

# 1. Parse RIEPILOGO
s_rep = wb["RIEPILOGO"]
for r in range(3, s_rep.max_row + 1):
    n = s_rep.cell(r, 1).value
    spec = s_rep.cell(r, 2).value
    if not spec or str(spec).strip().startswith("TOTALE"):
        continue
    stat = s_rep.cell(r, 3).value or 0
    reg = s_rep.cell(r, 4).value or 0
    altri = s_rep.cell(r, 5).value or 0
    tot = s_rep.cell(r, 6).value or 0
    summary_records.append({
        "n": n,
        "disciplina_mur": str(spec).strip(),
        "statali_2022": int(stat),
        "regionali_2022": int(reg),
        "altri_2022": int(altri),
        "totale_2022": int(tot)
    })

df_summary = pd.DataFrame(summary_records)
print(f"Summary sheet parsed: {len(df_summary)} specialties. Total contracts 2022: {df_summary['totale_2022'].sum()}")

# 2. Parse individual sheets
for sname in wb.sheetnames:
    if sname == "RIEPILOGO":
        continue
    sheet = wb[sname]
    spec_full = sheet.cell(1, 1).value or sname
    spec_full = str(spec_full).strip()
    
    for r in range(3, sheet.max_row + 1):
        ateneo = sheet.cell(r, 2).value
        if not ateneo or str(ateneo).strip().startswith("TOTALE"):
            continue
        ateneo = str(ateneo).strip()
        reg_sede = uni_to_region.get(ateneo, "Altro")
        
        stat = sheet.cell(r, 3).value or 0
        reg = sheet.cell(r, 4).value or 0
        note_reg = sheet.cell(r, 5).value or ""
        altri = sheet.cell(r, 6).value or 0
        
        all_regional_records.append({
            "disciplina_mur": spec_full,
            "sheet_name": sname,
            "ateneo": ateneo,
            "regione_sede": reg_sede,
            "posti_statali": int(stat),
            "posti_regionali": int(reg),
            "note_regionali": str(note_reg).strip(),
            "posti_altri": int(altri)
        })

df_details = pd.DataFrame(all_regional_records)
print(f"Individual sheets parsed: {len(df_details)} university-specialty offerings across Italy")
print(f"Total statali: {df_details['posti_statali'].sum()}, total regionali: {df_details['posti_regionali'].sum()}")
