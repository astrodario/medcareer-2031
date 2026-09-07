import openpyxl
import re

wb = openpyxl.load_workbook("data/raw/mur_ssm/riparto_SSM_22_23.xlsx", data_only=True)

regioni_finanziatrici = set()
for sname in wb.sheetnames:
    if sname == "RIEPILOGO":
        continue
    sheet = wb[sname]
    for r in range(3, sheet.max_row + 1):
        note = sheet.cell(r, 5).value
        if note and str(note).strip():
            regioni_finanziatrici.add(str(note).strip())

print(f"Unique regional notes count: {len(regioni_finanziatrici)}")
for n in sorted(regioni_finanziatrici)[:25]:
    print("  Note:", repr(n))
