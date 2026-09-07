import openpyxl

wb = openpyxl.load_workbook("data/raw/mur_ssm/riparto_SSM_22_23.xlsx", data_only=True)

universities = set()
for sheet_name in wb.sheetnames:
    if sheet_name == "RIEPILOGO":
        continue
    sheet = wb[sheet_name]
    # Header is typically row 2, data starts row 3
    for r in range(3, sheet.max_row + 1):
        ateneo = sheet.cell(r, 2).value
        if ateneo and str(ateneo).strip() and not str(ateneo).strip().startswith("TOTALE"):
            universities.add(str(ateneo).strip())

print(f"Total unique universities found: {len(universities)}")
for u in sorted(universities):
    print(f"  - '{u}'")
